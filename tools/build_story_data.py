# -*- coding: utf-8 -*-
"""从《2.0 物品清单》和《2.0 买家详细故事文档》生成 js/v2/storydata.js。

- 线索物品：名称、所属买家、类型（核心/线索一/线索二）、发现方式、关键词。
- 万物通结果：按买家整理成「关键词/组合 -> 结果类型 + 内容」。
"""

import os
import sys

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ITEMS_DOC = os.path.join(ROOT, "再拍最后一箱 2.0 物品清单.docx")
STORY_DOC = os.path.join(ROOT, "再拍最后一箱 2.0 买家详细故事文档.docx")
OUTPUT = os.path.join(ROOT, "js", "v2", "storydata.js")

# 线索物品的价值清单（物品清单文档没写，按物品性质给的基础价）。
CLUE_VALUES = {
    "手织毛衣": 140,
    "未寄出的邮寄单": 60,
    "毛线记账本": 190,
    "旧账本": 420,
    "手写转账单": 90,
    "无名感谢信": 60,
    "旧行李箱": 380,
    "未寄出的报警回执": 80,
    "当票存根": 70,
    "超期寄存箱": 260,
    "旧值班表": 90,
    "手写入库单": 80,
    "十二件套银餐具": 1600,
    "拍卖行成交单": 120,
    "修复店工单": 90,
    "旧录音带": 260,
    "手写节目单": 110,
    "电台工牌": 180
}

# 所有物品都要清洁和放大镜；只有「拆开夹页」这类需要拆装的才加电路修复。
DISCOVERY_OPS = {
    "清洁和整理": ["clean", "magnifier"],
    "清洁和校对": ["clean", "magnifier"],
    "标签检查": ["clean", "magnifier"],
    "检查标签": ["clean", "magnifier"],
    "放大镜检查": ["clean", "magnifier"],
    "放大镜": ["clean", "magnifier"],
    "纸张修复": ["clean", "magnifier"],
    "拆开夹页": ["clean", "circuit", "magnifier"]
}

# 买家顺序：物品清单和故事文档都是这个顺序。
BUYER_IDS = [
    "paper_crane",
    "blue_ledger",
    "old_fox",
    "seven_warehouse",
    "velvet_broker",
    "midnight_radio"
]

TYPE_MAP = {"核心": "core", "线索一": "clue1", "线索二": "clue2"}


def read_clue_items():
    document = docx.Document(ITEMS_DOC)
    table = document.tables[1]
    rows = []
    for index, row in enumerate(table.rows):
        if index == 0:
            continue
        cells = [cell.text.strip() for cell in row.cells]
        if len(cells) < 5 or not cells[0]:
            continue
        rows.append(
            {
                "name": cells[0],
                "buyer": cells[1],
                "type": TYPE_MAP.get(cells[2], "clue2"),
                "discovery": cells[3],
                "keywords": [part.strip() for part in cells[4].replace("、", ",").split(",") if part.strip()]
            }
        )
    return rows


def read_story_results():
    """故事文档里每个买家的物品表和结果表按顺序排列。"""
    document = docx.Document(STORY_DOC)
    tables = document.tables
    buyers = []
    cursor = 0
    for buyer_id in BUYER_IDS:
        # 每个买家：性格表(6x2) + 物品表(4x4) + 3 张结果表(4x3) + 选项表
        while cursor < len(tables) and not (
            len(tables[cursor].rows) == 4
            and len(tables[cursor].columns) == 4
            and tables[cursor].rows[0].cells[0].text.strip() == "物品"
        ):
            cursor += 1
        items_table = tables[cursor]
        cursor += 1
        results = {}
        combos = {}
        keywords = []
        for index, row in enumerate(items_table.rows):
            if index == 0:
                continue
            cells = [cell.text.strip() for cell in row.cells]
            if not cells[0]:
                continue
            keywords.extend(
                part.strip()
                for part in cells[2].replace("、", ",").split(",")
                if part.strip()
            )
        normalized = {key.replace(" ", ""): key for key in keywords}
        for _ in range(3):
            if cursor >= len(tables):
                break
            table = tables[cursor]
            cursor += 1
            header = [cell.text.strip() for cell in table.rows[0].cells]
            if header[:3] != ["搜索组合", "结果类型", "内容"]:
                break
            for row in table.rows[1:]:
                key = row.cells[0].text.strip()
                kind = row.cells[1].text.strip()
                text = row.cells[2].text.strip()
                if not key:
                    continue
                # 判断这一行用到哪几个关键词：1 个是单条结果，多个是组合检索。
                flat = key.replace(" ", "").replace("|", "")
                matched = [origin for norm, origin in normalized.items() if norm and norm in flat]
                entry = {"type": kind, "text": text}
                if len(matched) >= 2:
                    combos["|".join(sorted(matched))] = entry
                elif len(matched) == 1:
                    results.setdefault(matched[0], entry)
        buyer = {"id": buyer_id, "items": [], "results": results, "combos": combos}
        for index, row in enumerate(items_table.rows):
            if index == 0:
                continue
            cells = [cell.text.strip() for cell in row.cells]
            if not cells[0]:
                continue
            buyer["items"].append(
                {
                    "name": cells[0],
                    "type": TYPE_MAP.get(cells[1], "clue2"),
                    "keywords": [part.strip() for part in cells[2].replace("、", ",").split(",") if part.strip()],
                    "discovery": cells[3]
                }
            )
        buyers.append(buyer)
    return buyers


def iter_blocks(document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


# 选项语气 -> 信任变化与价格影响。
def option_effect(reaction):
    if "信任提高" in reaction:
        trust = 6
    elif "信任略升" in reaction:
        trust = 3
    elif "信任略降" in reaction:
        trust = -3
    elif "信任下降" in reaction:
        trust = -8
    else:
        trust = 0
    return {
        "trust": trust,
        "priceBoost": 0.1 if "成交价提高" in reaction or "收益提高" in reaction or "报价提高" in reaction else 0,
        "mayLeave": "立刻离店" in reaction or "直接离开" in reaction,
        "pause": "故事暂停" in reaction
    }


def read_special_dialogues():
    """每位买家的「5. 特殊会话和玩家选项」：开场台词 + 三个选项。"""
    document = docx.Document(STORY_DOC)
    specials = {}
    buyer_index = -1
    collecting = False
    lines = []
    options = []
    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text.endswith("5. 特殊会话和玩家选项"):
                buyer_index += 1
                collecting = True
                lines = []
                options = []
                continue
            if collecting and text.endswith("6. 记事本串联"):
                if 0 <= buyer_index < len(BUYER_IDS):
                    specials[BUYER_IDS[buyer_index]] = {"lines": lines, "options": options}
                collecting = False
                continue
            if not collecting or not text:
                continue
            if text.startswith("买家："):
                lines.append(text.replace("买家：", "", 1).strip())
            elif text.startswith("查到"):
                # 「查到 XX 后：」记录下一条台词属于哪件线索物品。
                lines.append("@" + text.replace("后：", "").strip())
        else:
            if not collecting:
                continue
            header = [cell.text.strip() for cell in block.rows[0].cells]
            if header[:3] != ["选项", "买家反应", "后续"]:
                continue
            for row in block.rows[1:]:
                cells = [cell.text.strip() for cell in row.cells]
                if not cells[0]:
                    continue
                effect = option_effect(cells[1])
                options.append(
                    {
                        "text": cells[0],
                        "reaction": cells[1],
                        "followUp": cells[2],
                        "trust": effect["trust"],
                        "priceBoost": effect["priceBoost"],
                        "mayLeave": effect["mayLeave"],
                        "pause": effect["pause"]
                    }
                )
    return specials


def js_string(value):
    return '"' + value.replace("\\", "\\\\").replace('"', '\\"') + '"'


def main():
    clue_items = read_clue_items()
    buyers = read_story_results()
    specials = read_special_dialogues()
    lines = [
        "// 本文件由 tools/build_story_data.py 从《2.0 物品清单》和《2.0 买家详细故事文档》生成。",
        "// 需要改线索物品、关键词或万物通结果时改文档，再重新跑一次脚本。",
        "",
        "// 线索物品：每件只会开出一个，开出后解锁对应买家故事。",
        "export const CLUE_ITEMS = ["
    ]
    for item in clue_items:
        ops = DISCOVERY_OPS.get(item["discovery"], ["magnifier"])
        lines.append("  {")
        lines.append("    id: %s," % js_string("clue_" + item["name"]))
        lines.append("    name: %s," % js_string(item["name"]))
        lines.append("    buyer: %s," % js_string(item["buyer"]))
        lines.append("    type: %s," % js_string(item["type"]))
        lines.append("    discovery: %s," % js_string(item["discovery"]))
        lines.append("    operations: [%s]," % ", ".join(js_string(op) for op in ops))
        lines.append("    keywords: [%s]," % ", ".join(js_string(k) for k in item["keywords"]))
        value = CLUE_VALUES.get(item["name"], 200)
        category = "luxury" if value >= 1000 else "common" if value >= 150 else "junk"
        lines.append("    baseValue: %d," % value)
        lines.append("    category: %s" % js_string(category))
        lines.append("  },")
    lines.append("];")
    lines.append("")
    lines.append("// 每篇故事的万物通检索结果，键是关键词或关键词组合。")
    lines.append("export const STORY_SEARCH = {")
    for buyer in buyers:
        lines.append("  %s: {" % js_string(buyer["id"]))
        lines.append("    results: {")
        for key, value in buyer["results"].items():
            lines.append(
                "      %s: { type: %s, text: %s },"
                % (js_string(key), js_string(value["type"]), js_string(value["text"]))
            )
        lines.append("    },")
        lines.append("    combos: {")
        for key, value in buyer["combos"].items():
            lines.append(
                "      %s: { type: %s, text: %s },"
                % (js_string(key), js_string(value["type"]), js_string(value["text"]))
            )
        lines.append("    }")
        lines.append("  },")
    lines.append("};")
    lines.append("")
    lines.append("// 每篇故事的特殊会话：线索物品到手后买家的台词与三个玩家选项。")
    lines.append("export const STORY_SPECIAL = {")
    for buyer_id, special in specials.items():
        lines.append("  %s: {" % js_string(buyer_id))
        lines.append("    lines: [%s]," % ", ".join(js_string(line) for line in special["lines"]))
        lines.append("    options: [")
        for option in special["options"]:
            lines.append("      {")
            lines.append("        text: %s," % js_string(option["text"]))
            lines.append("        reaction: %s," % js_string(option["reaction"]))
            lines.append("        followUp: %s," % js_string(option["followUp"]))
            lines.append("        trust: %d," % option["trust"])
            lines.append("        priceBoost: %s," % ("true" if option["priceBoost"] else "false"))
            lines.append("        mayLeave: %s," % ("true" if option["mayLeave"] else "false"))
            lines.append("        pause: %s" % ("true" if option["pause"] else "false"))
            lines.append("      },")
        lines.append("    ]")
        lines.append("  },")
    lines.append("};")
    lines.append("")
    with open(OUTPUT, "w", encoding="utf-8") as handle:
        handle.write("\n".join(lines))
    print("wrote", OUTPUT)
    print("clue items:", len(clue_items), "buyers:", len(buyers))
    return 0


if __name__ == "__main__":
    sys.exit(main())
