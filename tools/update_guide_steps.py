# -*- coding: utf-8 -*-
"""按当前实现同步新手引导策划案的第 1 天步骤表。"""

import copy
import os
import sys

import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GUIDE = os.path.join(ROOT, "再拍最后一箱 2.0 新手引导策划案.docx")

STEPS = [
    ["步骤", "触发", "小黑位置", "小黑台词或内容", "玩家完成条件"],
    ["1 看箱子", "开屏后", "当前箱子下方", "先看来源、重量、封条和寄存时间，再看右边的估值区间。", "点击「继续」"],
    ["2 出价", "上一步完成", "加价按钮左侧", "点这里出价。", "点击报价"],
    ["3 竞争对手", "上一步完成", "竞争对手组下方", "他们是今天的竞争对手，跟价很凶。继续出价把箱子拿下来。", "赢得当前箱子"],
    ["4 带回家", "上一步完成", "带回家卡片下方", "第一只箱子先带回家，另外两条路线暂时锁着。带回家的箱子会放在货架下层。", "选择带回家"],
    ["5 打开箱子", "回到店铺", "下层箱子格下方", "点下层箱子，箱内的物品会自动进入上层物品格。", "打开箱子"],
    ["6 拖到桌面", "箱子打开后", "上层物品格下方", "把上层物品拖到桌面，检视台才会打开。", "把物品拖到桌面"],
    ["7 检视操作", "检视台打开", "检视工具旁", "按物品性质完成全部常规操作，污渍要擦掉，真实信息要核对。", "完成全部常规操作"],
    ["8 回到货架", "检视完成", "上层物品格下方", "检视完成，物品回到货架。点开物品详情，再点出售。", "打开出售页面"],
    ["9 查看趋势", "出售页面打开", "日历下方", "先别急着定价，看看今天的市场趋势。", "打开日历"],
    ["10 出售定价", "查看趋势后", "出售页面", "记住今天的趋势，回到出售页面选两个标签并定价。", "完成定价并上架"],
    ["11 买家", "买家到店", "买家气泡下方", "买家会关心不同的信息，回答会影响信任和成交。", "完成一次成交"],
    ["12 返回拍卖场", "成交结束", "拍卖锤上方", "点一下桌面上的拍卖锤，就能回到拍卖场。", "点击拍卖锤"],
    ["13 自主竞拍", "回到拍卖场", "加价按钮左侧", "试着自己拍下一只箱子吧。你已经可以独立经营了！", "竞拍第二只箱子"],
]


def style_like(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def rebuild(table, rows):
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for row_index, values in enumerate(rows):
        if row_index >= len(table.rows):
            table.add_row()
        for col_index, value in enumerate(values):
            cell = table.rows[row_index].cells[col_index]
            style_like(cell.paragraphs[0], value)
            for extra in cell.paragraphs[1:]:
                extra._element.getparent().remove(extra._element)


def main():
    document = docx.Document(GUIDE)
    rebuild(document.tables[6], STEPS)
    document.save(GUIDE)
    print("updated steps: %d rows" % len(STEPS))


if __name__ == "__main__":
    sys.exit(main())
