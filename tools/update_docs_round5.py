# -*- coding: utf-8 -*-
"""同步：桌面物品方案、弹层层级、开场与自我介绍居中、日历关闭保留出售页面。"""

import copy
import os
import sys

import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GUIDE = os.path.join(ROOT, "再拍最后一箱 2.0 新手引导策划案.docx")
SHOP = os.path.join(ROOT, "再拍最后一箱 2.0 店铺界面策划案.docx")


def find(paragraphs, keyword, start=0):
    for index in range(start, len(paragraphs)):
        if keyword in paragraphs[index].text:
            return index
    return -1


def set_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return paragraph
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""
    return paragraph


def insert_after(anchor, text):
    new = anchor.insert_paragraph_before("", style=anchor.style)
    anchor._element.addnext(new._element)
    run = new.add_run(text)
    rpr = anchor._element.find(qn("w:rPr"))
    if rpr is not None:
        run._element.insert(0, copy.deepcopy(rpr))
    return new


def main():
    shop = docx.Document(SHOP)
    paragraphs = shop.paragraphs

    # 桌面物品：悬停显示名称和提示，点击打开物品信息窗口。
    index = find(paragraphs, "买家到店开始谈价时，待谈物品会摆在桌面上")
    if index >= 0:
        set_text(
            paragraphs[index],
            "买家到店开始谈价时，待谈物品会自动摆在桌面正中央。",
        )
        anchor = paragraphs[index]
        anchor = insert_after(
            anchor,
            "鼠标悬停在桌面物品上时显示物品名称和提示框「点击查看物品详细信息」。",
        )
        anchor = insert_after(
            anchor,
            "点击桌面物品打开物品信息窗口，内容与检视记录一致：物品状态、各项检视操作结果、"
            "已发现问题、万物通参考（参考售价、当前趋势、标签标注年代）和买入价。",
        )
        insert_after(
            anchor,
            "物品信息窗口只是随身查看，不影响买家对话和交易进度；关闭后回到买家交易界面。",
        )

    # 界面层级：弹层要盖住买家气泡和选项。
    index = find(paragraphs, "界面层级：场景在最下")
    if index >= 0:
        set_text(
            paragraphs[index],
            "界面层级从下到上依次是：场景 → 买家对话气泡和三个选项 → 所有弹层（物品详情、物品信息窗口、"
            "出售页面、商城、检视台、日历、记事本、万物通）→ 小黑提示和右上角提示条。弹层必须盖住买家气泡，"
            "不会出现弹窗被气泡或界面信息压住的情况。",
        )

    index = find(paragraphs, "买家开始交易时，对应物品自动出现在桌面上")
    if index >= 0:
        set_text(paragraphs[index], "• 买家开始交易时，对应物品自动出现在桌面正中央。")
    index = find(paragraphs, "• 鼠标悬停桌面的物品时，显示名称、买入价、主要风险和参考售价等简要信息。")
    if index >= 0:
        set_text(paragraphs[index], "• 鼠标悬停桌面物品时显示物品名称和提示框「点击查看物品详细信息」。")
    index = find(paragraphs, "• 点击桌面物品，弹出物品详情窗口。")
    if index >= 0:
        set_text(
            paragraphs[index],
            "• 点击桌面物品打开物品信息窗口，显示物品状态、检视记录、已发现问题、万物通参考和买入价。",
        )

    index = find(paragraphs, "• 店铺升级必须同时改变画面和经营效率。")
    if index >= 0:
        anchor = paragraphs[index]
        anchor = insert_after(anchor, "• 买家交易时物品出现在桌面正中央，悬停显示名称和「点击查看物品详细信息」提示。")
        anchor = insert_after(anchor, "• 所有弹层都盖在买家气泡和三个选项之上，弹窗不会被气泡遮住。")
        insert_after(anchor, "• 从出售页面打开日历再关闭，出售页面和已设置的价格、标签、展示天数全部保留。")
    shop.save(SHOP)
    print("updated: 店铺界面策划案")

    guide = docx.Document(GUIDE)
    paragraphs = guide.paragraphs

    index = find(paragraphs, "开场第一段小黑对话出现在画面正中央")
    if index >= 0:
        insert_after(
            paragraphs[index],
            "选择「你谁啊」后的自我介绍对话同样在画面正中央，小黑站在对话气泡左侧，背景保持压暗，"
            "玩家点「开始竞拍」才进入拍卖。",
        )

    index = find(paragraphs, "5. 玩家打开日历看完趋势，小黑提示回到出售页面选两个标签并定价。")
    if index >= 0:
        insert_after(
            paragraphs[index],
            "关闭日历后出售页面不会跟着一起关掉，已经调好的价格、标签和展示天数全部保留，"
            "玩家可以直接继续定价上架。",
        )

    index = find(paragraphs, "• 悬停桌面物品时显示简要信息。")
    if index >= 0:
        set_text(paragraphs[index], "• 悬停桌面物品时显示物品名称和提示框「点击查看物品详细信息」。")
    index = find(paragraphs, "• 点击物品弹出详细物品详情窗口。")
    if index >= 0:
        set_text(paragraphs[index], "• 点击桌面物品弹出物品信息窗口，显示物品状态和检视详细记录。")
    index = find(paragraphs, "买家到店后，对应物品会出现在桌面上。小黑在买家步骤中说明：")
    if index >= 0:
        set_text(paragraphs[index], "买家到店后，对应物品会出现在桌面正中央。小黑在买家步骤中说明：")

    index = find(paragraphs, "• 同一时间只有一个教学弹窗。")
    if index >= 0:
        anchor = paragraphs[index]
        anchor = insert_after(anchor, "• 开场对话和自我介绍对话都出现在画面正中央。")
        anchor = insert_after(anchor, "• 教学期第一次打开日历再关闭，出售页面和已经填好的价格、标签不会丢失。")
        insert_after(anchor, "• 教学期点击桌面待谈物品可以打开物品信息窗口，关闭后继续买家对话。")
    guide.save(GUIDE)
    print("updated: 新手引导策划案")


if __name__ == "__main__":
    sys.exit(main())
