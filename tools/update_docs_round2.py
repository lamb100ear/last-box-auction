# -*- coding: utf-8 -*-
"""同步小黑引导、弹层层级、检视玩法和万物通窗口的改动。

新手引导策划案：小黑跟随指引目标定位、严格按顺序推进、第 1 天步骤表和弹窗层级。
店铺界面策划案：检视玩法按物品性质配对、电路修复的额外情况、弹层关闭按钮和万物通窗口。
"""

import copy
import os
import sys

import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GUIDE = os.path.join(ROOT, "再拍最后一箱 2.0 新手引导策划案.docx")
SHOP = os.path.join(ROOT, "再拍最后一箱 2.0 店铺界面策划案.docx")


def style_like(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def find(paragraphs, keyword, start=0):
    for index in range(start, len(paragraphs)):
        if keyword in paragraphs[index].text:
            return index
    raise SystemExit("找不到段落：%s" % keyword)


def insert_after(anchor, text):
    new = anchor.insert_paragraph_before("", style=anchor.style)
    anchor._element.addnext(new._element)
    run = new.add_run(text)
    rpr = anchor._element.find(qn("w:rPr"))
    if rpr is not None:
        run._element.insert(0, copy.deepcopy(rpr))
    return new


def rebuild_table(table, rows):
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for row_index, values in enumerate(rows):
        if row_index >= len(table.rows):
            table.add_row()
        for col_index, value in enumerate(values):
            cell = table.rows[row_index].cells[col_index]
            paragraph = cell.paragraphs[0]
            style_like(paragraph, value)
            for extra in cell.paragraphs[1:]:
                extra._element.getparent().remove(extra._element)


DAY1_STEPS = [
    ["步骤", "触发", "小黑位置", "小黑台词或内容", "玩家完成条件"],
    ["1 看箱子", "开屏后", "当前箱子下方", "先看来源、重量、封条和寄存时间，再看右边的估值区间。", "点击「继续」"],
    ["2 出价", "上一步完成", "加价按钮左侧", "点这里出价。", "点击报价"],
    ["3 竞争对手", "上一步完成", "竞争对手组下方", "他们是今天的竞争对手，跟价很凶。", "出现跟价，或点击「继续」"],
    ["4 赢得第一只箱子", "上一步完成", "带回家卡片下方", "第一只箱子先带回家，学会处理流程，另外两条路线暂时锁着。", "进入处置界面"],
    ["5 带回家", "进入处置界面", "带回家卡片下方", "带回家的箱子放在货架下层。", "选择带回家"],
    ["6 打开箱子", "回到店铺", "下层箱子格下方", "点下层箱子，箱内的物品会自动进入上层物品格。", "打开箱子"],
    ["7 拖到桌面", "箱子打开后", "上层物品格下方", "把上层物品拖到桌面，检视台才会打开。", "把物品拖到桌面"],
    ["8 检视", "检视台打开", "检视台右侧", "先用清洁刷把污渍擦掉，再用放大镜核对真实信息。", "完成一项常规操作"],
    ["9 看市场趋势", "上一步完成", "日历下方", "先看今天的市场趋势，再决定价格。", "打开日历"],
    ["10 出售定价", "查看趋势后", "橱窗下方", "出售要选两个标签，价格定在参考区间内。", "完成一次定价并上架"],
    ["11 买家", "买家到店", "买家气泡下方", "买家会关心不同的信息，回答会影响信任和成交。", "完成一次成交"],
    ["12 返回拍卖场", "成交结束", "拍卖锤上方", "点一下桌面上的拍卖锤，就能回到拍卖场。", "点击拍卖锤"],
    ["13 自主竞拍", "回到拍卖场", "加价按钮左侧", "试着自己拍下一只箱子吧。你已经可以独立经营了！", "竞拍第二只箱子"],
]

POPUP_LAYERS = [
    ["层级", "内容", "是否阻断"],
    ["1", "场景和界面（拍卖场、处置台、店铺）", "否"],
    ["2", "常驻状态栏（余额、麻烦值、信誉值）和日历缩略", "否"],
    ["3", "高亮轮廓、目标箭头和步骤提示卡", "否"],
    ["4", "小黑头像和对话气泡", "只有气泡里的按钮拦截点击"],
    ["5", "店铺弹层：物品详情、出售页面、商城、检视台、日历、记事本、万物通", "是"],
    ["6", "开场选择、贷款合同和系统错误", "是"],
]


def update_guide():
    document = docx.Document(GUIDE)
    paragraphs = document.paragraphs

    # 小黑位置：改为跟随当前指引目标。
    start = find(paragraphs, "小黑默认出现在场景左侧")
    texts = [
        "• 小黑和气泡作为整体，跟随当前步骤的目标 UI 移动，不再固定在左下角。",
        "• 定位顺序为：目标下方、目标左侧、目标上方、目标右侧，取第一个不越界的位置。",
        "• 小黑始终紧贴目标控件，方便玩家一边看提示一边操作。",
        "• 气泡里的按钮始终可以点击，头像和文字不拦截鼠标操作。",
        "• 气泡最多占画面宽度的 60%，任何情况下都不覆盖当前目标控件。",
        "• 气泡不提供「收起」按钮，避免误关闭引导。",
    ]
    style_like(document.paragraphs[start], texts[0])
    for index in range(start + 1, start + 9):
        element = document.paragraphs[index]._element
        element.getparent().remove(element)
    anchor = document.paragraphs[start]
    for text in texts[1:]:
        anchor = insert_after(anchor, text)

    # 严格按顺序推进。
    check = find(document.paragraphs, "每一步按以下顺序检查")
    anchor = document.paragraphs[check]
    for text in [
        "步骤推进规则：每个步骤都有唯一的完成条件，只有当前步骤的完成条件满足后才会进入下一步。"
        "其他动作不会让流程跳步，避免只点一次加价就跳到处置界面。",
        "如果某一步的条件暂时无法触发（例如没有竞争对手跟价），玩家可以点气泡里的「继续」手动进入下一步，"
        "不会卡住流程。",
    ]:
        anchor = insert_after(anchor, text)

    rebuild_table(document.tables[4], POPUP_LAYERS)
    rebuild_table(document.tables[6], DAY1_STEPS)
    document.save(GUIDE)
    print("updated: 新手引导策划案")


def update_shop():
    document = docx.Document(SHOP)

    # 检视玩法：按物品性质配对操作 + 电路修复的额外情况。
    anchor = document.paragraphs[find(document.paragraphs, "补充标签不再单独出现在检视台")]
    for text in [
        "操作按物品性质配对，不是每件物品都有三项：档案、票据、纸张和织物类只做清洁和放大镜检查，"
        "不做拆装；带电路或机械结构的物品才有拆解操作；纯机械物品（表、罗盘、剪刀、工具）"
        "收尾是「完成装机」，只有带电物品才做「通电测试」。",
        "电路修复揭开外壳盖板后，会随机出现一种额外情况：",
        "1. 零件损坏：拆下的零件已经损坏，无法装回，物品价值下降，本次评价最高只能到合格。",
        "2. 内部污渍：外壳内部还有一层污渍，需要用清洁布按住拖动擦掉，否则无法收尾。",
        "3. 安装顺序错误：零件必须按照提示的顺序装回，顺序错了会留下缝隙并记一次失误。",
    ]:
        anchor = insert_after(anchor, text)

    # 弹层层级与关闭按钮。
    layer_anchor = document.paragraphs[find(document.paragraphs, "检视台是桌面上的弹出层")]
    for text in [
        "所有弹层（物品详情、出售页面、商城、检视台、日历、记事本、万物通）都盖在常驻状态栏和日历缩略之上，"
        "弹层打开时状态栏被遮住，不会出现弹层被界面信息压住的情况。",
        "每个弹层右上角都有统一的关闭按钮「×」，点击后回到店铺主界面。",
        "记事本窗口有三个页签：买家故事、图鉴、店铺升级，切换页签时三个入口始终可见。",
        "点击「用万物通查询」后，左侧打开万物通检索窗口，样式参考电脑上的搜索页面："
        "顶部是检索栏和关键词，下面是新闻、档案、帖子和矛盾记录卡片。"
        "记事本在右侧、万物通在左侧，两个窗口同时打开也不会互相遮挡。",
    ]:
        layer_anchor = insert_after(layer_anchor, text)

    document.save(SHOP)
    print("updated: 店铺界面策划案")


def main():
    update_guide()
    update_shop()


if __name__ == "__main__":
    sys.exit(main())
