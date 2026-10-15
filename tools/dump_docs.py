# -*- coding: utf-8 -*-
"""列出策划案文档的段落索引，供人工定位锚点。输出写到 tools/_dump.txt。"""

import os
import sys

import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def main():
    lines = []
    for name in sys.argv[1:]:
        doc = docx.Document(os.path.join(ROOT, name))
        lines.append("=" * 30 + " " + name)
        for index, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if text:
                lines.append(f"{index} | {paragraph.style.name} | {text[:130]}")
        lines.append(f"--tables: {len(doc.tables)}")
        for table_index, table in enumerate(doc.tables):
            lines.append(f"--table {table_index}: {len(table.rows)}x{len(table.columns)}")
            for row in table.rows[:16]:
                lines.append("    " + " | ".join(cell.text.strip()[:26] for cell in row.cells))
    output = os.path.join(ROOT, "tools", "_dump.txt")
    with open(output, "w", encoding="utf-8") as handle:
        handle.write("\n".join(lines))
    print("wrote", output)


if __name__ == "__main__":
    sys.exit(main())
