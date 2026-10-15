# -*- coding: utf-8 -*-
"""同步：取消第二天引导、首次上架保证买家、检视台按键高亮。"""

import copy
import os
import sys

import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GUIDE = os.path.join(ROOT, "再拍最后一箱 2.0 新手引导策划案.docx")
SHOP = os.path.join(ROOT, "再拍最后一箱 2.0 店铺界面策划案.docx")


def style_like(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def find(paragraphs, keyword, start=0):
    for index in range(start, len(paragraphs)):
        if keyword in paragraphs[index].text:
            return index
    return -1


def insert_after(anchor, text):
    new = anchor.insert_paragraph_before("", style=anchor.style)
    anchor._element.addnext(new._element)
    run = new.add_run(text)
    rpr = anchor._element.find(qn("w:rPr"))
    if rpr is not None:
        run._element.insert(0, copy.deepcopy(rpr))
    return new


def rebuild(table, rows):
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for row_index, values in enumerate(rows):
        if row_index >= len(table.rows):
            table.add_row()
        for col_index, value in enumerate(values):
            cell = table.rows[row_index].cells[col_index]
            style_like(cell.paragraphs[0], value)
            for extra in cell.paragraphs[1:]:
                extra._element.getparent().remove(extra._element)


def main():
    guide = docx.Document(GUIDE)
    # 第二天不再有引导。
    for paragraph in guide.paragraphs:
        if paragraph.text.strip() == "今天市场变了。别急着卖，先看日历。":
            style_like(paragraph, "第 2 天不再出现任何小黑引导，玩家自行经营。")
            break
    rhythm = guide.tables[0]
    rows = [[cell.text for cell in row.cells] for row in rhythm.rows]
    for row in rows:
        if row[0].strip().startswith("第 2 天"):
            row[1] = "不再出现小黑引导，玩家自行经营"
        if row[0].strip().startswith("第 4 天"):
            row[1] = "高额箱和贷款分支，余额不足时小黑再次出现"
    rebuild(rhythm, rows)

    anchor = guide.paragraphs[find(guide.paragraphs, "日历步骤必须排在出售步骤之前")]
    for text in [
        "第一次出售的买家保证到店：完整教学期间第一次把物品放进橱窗后，买家会在几秒内到店，"
        "方便玩家理解上架→等待→到店→对话→成交的流程。",
        "教学结束后（或玩家选择跳过教学）恢复正常的买家到店概率，按店铺等级抽取到店时间。",
        "检视台里会用高亮描边标出下一步该点的按键：先亮清洁刷，清洁完成亮确认按钮，"
        "再亮放大镜，找齐真实信息后亮确认按钮，全部完成后亮完成检视。",
    ]:
        anchor = insert_after(anchor, text)
    guide.save(GUIDE)
    print("updated: 新手引导策划案")

    shop = docx.Document(SHOP)
    anchor = shop.paragraphs[find(shop.paragraphs, "电路修复揭开外壳盖板后")]
    for text in [
        "内部污渍擦掉之后就可以继续拖零件和完成收尾，不会卡住流程。",
        "出售页面的价格滑杆可以按住拖动，价格提示会随拖动即时更新，不需要重新打开页面。",
    ]:
        anchor = insert_after(anchor, text)
    shop.save(SHOP)
    print("updated: 店铺界面策划案")


if __name__ == "__main__":
    sys.exit(main())
