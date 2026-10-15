# -*- coding: utf-8 -*-
"""同步：新手引导买家步骤的成功与失败分支。"""

import copy
import os
import sys

import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GUIDE = os.path.join(ROOT, "再拍最后一箱 2.0 新手引导策划案.docx")


def find(paragraphs, keyword, start=0):
    for index in range(start, len(paragraphs)):
        if keyword in paragraphs[index].text:
            return index
    return -1


def set_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return paragraph
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""
    return paragraph


def insert_after(anchor, text, style=None):
    new = anchor.insert_paragraph_before("", style=style or anchor.style)
    anchor._element.addnext(new._element)
    run = new.add_run(text)
    if style is None:
        rpr = anchor._element.find(qn("w:rPr"))
        if rpr is not None:
            run._element.insert(0, copy.deepcopy(rpr))
    return new


BRANCH_LINES = [
    "买家步骤有两个分支，两条分支都要能继续教学：",
    "• 成功分支：成交或压价成交后，小黑按原台词提示玩家点击拍卖锤返回拍卖场。",
    "• 失败分支：编造被识破、买家嫌价格太高直接离店、信任不够谈崩、"
    "买家没等到回复取消交易，这四种情况都会切到失败分支。",
    "失败分支规则：",
    "1. 小黑换成安慰台词，例如「谈崩了？不要灰心。第一次谈不成很正常，重新上架就能再等一位买家。」。",
    "2. 失败提示出现在画面正中央，背景压暗，只有一个「继续」按钮。",
    "3. 玩家点「继续」后，教学直接进入下一步「返回拍卖场」，不会卡在买家步骤。",
    "4. 交易失败时物品已经回到货架，玩家可以重新上架，再等新的买家。",
    "5. 失败台词按原因区分：被识破、定价过高吓走买家、信任不够谈崩、买家取消交易。",
    "6. 失败分支只在教学期第一次遇到买家步骤时出现一次，之后的交易失败走正常的失败结算。",
]

VERIFY_LINES = [
    "• 教学期第一次交易失败时，小黑给出失败分支台词并继续推进教学，不会卡在买家步骤。",
    "• 失败分支的提示居中显示，只有一个「继续」按钮。",
]


def main():
    guide = docx.Document(GUIDE)
    paragraphs = guide.paragraphs

    # 1. 第 1 天顺序里的买家步骤
    index = find(paragraphs, "7. 买家到店后进入对话引导，成交后小黑提示点击拍卖锤返回拍卖场。")
    if index >= 0:
        set_text(
            paragraphs[index],
            "7. 买家到店后进入对话引导：成交或压价成交走成功分支，小黑提示点击拍卖锤返回拍卖场；"
            "交易失败走失败分支，小黑说一句「不要灰心」，玩家点继续后同样进入拍卖锤引导。",
        )
        anchor = paragraphs[index]
        for text in BRANCH_LINES:
            anchor = insert_after(anchor, text)

    # 2. 步骤表里的买家步骤
    if len(guide.tables) > 6:
        table = guide.tables[6]
        for row in table.rows:
            cells = row.cells
            if cells[0].text.strip().startswith("11"):
                set_text(cells[1].paragraphs[0], "买家到店")
                set_text(
                    cells[3].paragraphs[0],
                    "买家会关心不同的信息，回答会影响信任和成交。"
                    "失败时改为：谈崩了？不要灰心。物品已经回货架，重新上架就能再等一位买家。",
                )
                set_text(
                    cells[4].paragraphs[0],
                    "完成一次成交，或交易失败后点「继续」",
                )
            if cells[0].text.strip().startswith("12"):
                set_text(cells[1].paragraphs[0], "买家步骤结束")
            for extra in cells[0].paragraphs[1:]:
                extra._element.getparent().remove(extra._element)
            for index_cell in (1, 3, 4):
                for extra in cells[index_cell].paragraphs[1:]:
                    extra._element.getparent().remove(extra._element)

    # 3. 验收要点
    index = find(paragraphs, "• 教学期点击桌面待谈物品可以打开物品信息窗口，关闭后继续买家对话。")
    if index >= 0:
        anchor = paragraphs[index]
        for text in VERIFY_LINES:
            anchor = insert_after(anchor, text)

    guide.save(GUIDE)
    print("updated: 新手引导策划案")
    return 0


if __name__ == "__main__":
    sys.exit(main())
