# -*- coding: utf-8 -*-
"""同步：藏品与收纳盒、图鉴表达、伪造商城入口、麻烦值与威胁事件；并生成买家系统改版方案。"""

import copy
import os
import sys

import docx
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOP = os.path.join(ROOT, "再拍最后一箱 2.0 店铺界面策划案.docx")
MASTER = os.path.join(ROOT, "再拍最后一箱 2.0 总策划案.docx")
BUYER = os.path.join(ROOT, "再拍最后一箱 2.0 买家系统改版方案（待确认）.docx")


def find(paragraphs, keyword, start=0):
    for index in range(start, len(paragraphs)):
        if keyword in paragraphs[index].text:
            return index
    return -1


def set_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return paragraph
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""
    return paragraph


def insert_after(anchor, text, style=None):
    new = anchor.insert_paragraph_before("", style=style or anchor.style)
    anchor._element.addnext(new._element)
    run = new.add_run(text)
    if style is None:
        rpr = anchor._element.find(qn("w:rPr"))
        if rpr is not None:
            run._element.insert(0, copy.deepcopy(rpr))
    return new


def insert_many(anchor, lines):
    current = anchor
    for text in lines:
        current = insert_after(current, text)
    return current


def insert_heading(anchor, text, level=3):
    return insert_after(anchor, text, style=f"Heading {level}")


def insert_table(anchor, document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        table.style = document.tables[0].style
    except Exception:
        pass
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.rows[row_index].cells[col_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(value)
            run.font.size = Pt(9)
    anchor._element.addnext(table._tbl)
    return table


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


KEEPSAKE_RULES = [
    "藏品规则：",
    "• 藏品是带额外特殊属性的物品，属性直接写在物品的标签上，只有藏品可以自留。",
    "• 属性有两种生效方式：自留生效（放进收纳盒后长期生效）；售出生效（成功售出时一次性生效）。",
    "• 自留生效的属性：老主顾（买家初始信任 +5）、夹层收纳（每件物品保管费 -5）、"
    "展示加成（在线成交价 +5%）、现场人脉（现场出售收入 +10%）、转卖渠道（原封转卖报价 +8%）、"
    "档案线索（万物通查询多确认一条信息）、法律人脉（麻烦处理服务降价 ¥100）。",
    "• 售出生效的属性：口碑藏品（成功售出后信誉 +1）、溢价藏品（成功售出后额外获得成交价 5% 现金）。",
    "• 收纳盒最多放 4 件藏品，同类藏品每件还给对应品类 2% 的参考价加成。",
    "• 同一标签的自留藏品同时放进收纳盒只生效一次，重复摆放不会叠加，收纳盒里会直接标出"
    "「重复标签未生效」的提示。",
    "• 普通物品不能自留，物品详情窗口不显示自留按钮。",
    "• 收纳盒已满但玩家仍然点「自留」：弹出提示框「收纳盒已满，是否要进行替换？」。",
    "   - 选择「是」：打开收纳盒格子窗口，鼠标悬停每个格子显示「点击以替换藏品」，点击后当前物品"
    "进入该格，被替换下来的藏品回到货架，可以继续出售。",
    "   - 选择「否」：回到物品详情页面，物品仍然留在货架。",
]

TROUBLE_SECTION = [
    "麻烦值范围 0 到 10，常驻状态栏第二醒目位置。0 至 3 是安全区，4 至 6 是警告区，7 至 10 是危险区。",
    "麻烦值上升：",
    "• 现场开箱：当场开出奢侈品，每件 +1，单箱最多 +2；当场开出违禁品，每件 +2。",
    "• 带回家开箱：箱子里有违禁品时 +1。",
    "• 伪造被识破：低风险 +2、中风险 +2、高风险 +4。",
    "• 违禁品尾款被冻结：每次 +1。",
    "• 逾期未还款：每逾期一期 +1。",
    "• 威胁事件处理失败：+1。",
    "麻烦值下降：",
    "• 律师服务 -2；持有「法律人脉」藏品时服务降价 ¥100。",
    "• 数据清理 -1，并获得一次伪造护航。",
    "• 上交违禁品 -1。",
    "• 连续 3 天没有新的风险事件，每天 -1，最低 0。",
    "警察调查：麻烦值达到 7 时触发，按阶段推进。",
    "• 触发当天：便衣上门，要求查看货架和记录。配合检查则麻烦 -2、信誉 +3、罚款 ¥300 至 800；"
    "隐瞒则调查推进一级，并有 30% 概率当场查扣一件违禁品。",
    "• 调查阶段 1（警告）：罚款 ¥1,000，麻烦 -3，限 3 天内处理，否则进入阶段 2；"
    "律师服务可以直接解除这一阶段。",
    "• 调查阶段 2（查扣）：罚款 ¥2,500，随机查扣两件高价值物品，麻烦清零，信誉 -6；"
    "「保险文件袋」类藏品让查扣罚款减半。",
    "• 调查阶段 3（终局）：账户冻结，游戏结束。",
    "威胁事件：第 6 天起每隔 3 到 6 天随机触发一次；麻烦值达到 4 以后触发概率提高。",
    "• 保护服务可以抵消一次威胁事件，不付钱也没有后果。",
    "• 每次成功处理威胁，下一次威胁强度 +8%，最多 +40%。",
    "• 警察调查和威胁事件不会在同一天触发。",
]

THREAT_ROWS = [
    ["威胁事件", "表现", "处理方式"],
    ["前主上门", "有人声称箱子里的东西本来属于他", "归还物品（信誉 +2）或不归还（麻烦 +1，30% 概率损失该物品）"],
    ["中介催款", "中介要求补交通道费", "支付 ¥400 至 ¥900，或拒绝（麻烦 +1，下一次成交价 -10%）"],
    ["同行举报", "竞争对手举报你的来源", "支付 ¥600 至 ¥1,200，或接受检查（麻烦 +2）"],
    ["仓库盘查", "保管方要求补齐寄存登记", "补交保管费差额，或丢失一只未拆封的箱子"],
]


def update_shop():
    document = docx.Document(SHOP)
    paragraphs = document.paragraphs

    # 1. 藏品与收纳盒
    index = find(paragraphs, "收纳盒规则：")
    if index >= 0:
        anchor = paragraphs[index]
        for offset, text in enumerate(KEEPSAKE_RULES):
            if offset == 0:
                set_text(anchor, text)
                continue
            anchor = insert_after(anchor, text)
    index = find(paragraphs, "• 有属性的藏品放入收纳盒后长期生效。")
    if index >= 0:
        remove_paragraph(paragraphs[index])
        paragraphs = document.paragraphs
    # 旧的四条收纳盒说明已经被新的藏品规则覆盖，直接删掉避免重复。
    for keyword in [
        "• 四个格子满了以后可以替换藏品。",
        "• 被替换下来的藏品回到货架，可继续出售。",
        "• 收纳盒只收纳有属性的藏品，普通物品不能放入。"
    ]:
        index = find(paragraphs, keyword)
        if index >= 0:
            remove_paragraph(paragraphs[index])
            paragraphs = document.paragraphs

    # 2. 物品详情里的自留
    index = find(paragraphs, "• 自留后物品进入桌面收纳盒。")
    if index >= 0:
        anchor = set_text(paragraphs[index], "• 只有带属性的藏品才有自留按钮，点击后物品进入桌面收纳盒。")
        anchor = insert_after(
            anchor,
            "• 自留生效的属性和同类 2% 参考价加成立刻生效；售出生效的属性在该物品成功售出时结算。"
        )
        anchor = insert_after(
            anchor,
            "• 收纳盒 4 格已满时先确认是否替换，再从收纳盒格子里点选要替换的藏品。"
        )
        insert_after(anchor, "• 被替换下来的藏品回到货架，可以继续出售。")
    index = find(paragraphs, "• 收纳盒四个格子已满时，选择替换哪个藏品。")
    if index >= 0:
        set_text(paragraphs[index], "• 收纳盒已满时，提示框里选择「是」进入替换，选择「否」回到物品详情。")
    index = find(paragraphs, "• 有属性的物品自留后属性立即生效。")
    if index >= 0:
        set_text(paragraphs[index], "• 自留后收纳盒里会显示当前生效中的标签；重复标签只生效一次。")
    index = find(paragraphs, "• 被替换的藏品回到货架，可以继续出售。")
    if index >= 0:
        remove_paragraph(paragraphs[index])
        paragraphs = document.paragraphs
    index = find(paragraphs, "• 被替换下来的藏品回到货架")
    if index >= 0 and paragraphs[index].text.startswith("• 被替换下来的藏品回到货架"):
        # 保留第一条（自留流程里的说明），删掉重复的第二条。
        second = find(paragraphs, "• 被替换下来的藏品回到货架", index + 1)
        if second >= 0:
            remove_paragraph(paragraphs[second])
            paragraphs = document.paragraphs

    # 3. 伪造面板直接买材料
    index = find(paragraphs, "• 伪造页面不创建新窗口。")
    if index < 0:
        index = find(paragraphs, "补充标签不再单独出现在检视台")
    if index >= 0:
        anchor = paragraphs[index]
        insert_after(
            anchor,
            "伪造材料不足时，伪造面板直接提供「去电脑商城买材料」按钮，点开商城购买后回到检视台的伪造面板，"
            "材料和进度都会保留。"
        )

    # 4. 麻烦值、警察与威胁事件
    anchor = None
    for paragraph in paragraphs:
        if "第五次出售不默认被捕" in paragraph.text:
            anchor = paragraph
            break
    if anchor is not None:
        anchor = insert_heading(anchor, "麻烦值、警察与威胁事件", level=3)
        for text in TROUBLE_SECTION:
            anchor = insert_after(anchor, text, style="Normal")
        insert_table(anchor, document, THREAT_ROWS)

    # 5. 图鉴表达
    index = find(paragraphs, "图鉴分为：")
    if index >= 0:
        anchor = paragraphs[index]
        anchor = insert_after(anchor, "• 物品：每个已经见过的物品显示图片，没见过的显示剪影，"
                                      "除物品名称和成交次数外不显示其他文字。")
        anchor = insert_after(anchor, "• 箱子：按来源各占一格，已见过的显示箱子图片，没见过的显示剪影，"
                                      "格子里不写任何文字。")
        anchor = insert_after(anchor, "• 买家：保留头像、故事进度、关键词和故事残页。")
        insert_after(anchor, "• 物品图使用 美术/物品/by-id/<模板 ID>.png；箱子图使用 "
                             "美术/2.0/compendium/箱子_<来源 ID>.png；美术出图后覆盖同名文件即可。")
    index = find(paragraphs, "箱子和物品使用已见显示、未见解锁剪影。")
    if index >= 0:
        set_text(
            paragraphs[index],
            "箱子和物品使用已见显示图片、未见解锁剪影；买家档案使用独立头像、进度、关键词和故事残页。"
        )

    # 6. 验收要点
    index = find(paragraphs, "• 店铺升级必须同时改变画面和经营效率。")
    if index >= 0:
        anchor = paragraphs[index]
        anchor = insert_after(anchor, "• 只有带属性的藏品能自留，收纳盒最多 4 件，同标签只生效一次并有提示。")
        anchor = insert_after(anchor, "• 收纳盒满时自留会先问是否替换，替换下来的藏品回到货架。")
        anchor = insert_after(anchor, "• 现场开出奢侈品或违禁品时麻烦值会上升。")
        insert_after(anchor, "• 图鉴物品显示物品名称和成交次数，箱子格不显示任何文字。")

    document.save(SHOP)
    print("updated: 店铺界面策划案")


def update_master():
    document = docx.Document(MASTER)
    paragraphs = document.paragraphs
    index = find(paragraphs, "现金是唯一通用支付资源。")
    if index >= 0:
        insert_after(
            paragraphs[index],
            "麻烦值来自现场开箱的奢侈品和违禁品、伪造被识破、违禁品尾款冻结、逾期和威胁事件；"
            "麻烦值达到 7 会触发警察调查，第 6 天起会随机触发威胁事件。详细规则见"
            "《再拍最后一箱 2.0 店铺界面策划案》的「麻烦值、警察与威胁事件」。"
        )
    index = find(paragraphs, "伪造系统使用电脑商城购买材料")
    if index >= 0:
        insert_after(
            paragraphs[index],
            "借贷的压力来自利息而不是额度：额度按资金缺口给，利息随借款次数递增，"
            "第 4 天黑匣紧急竞拍垫资固定 300% 利息，还款分 3 期、共 9 天。详见"
            "《再拍最后一箱-数值平衡与资源换算表》第 13 章。"
        )
    document.save(MASTER)
    print("updated: 总策划案")


BUYER_CONTENT = [
    ("h1", "《再拍最后一箱！》2.0 买家系统改版方案（待确认）"),
    ("p", "本文档只提出买家系统的改版方案，暂不写入其它策划案，等待确认后再合并。"),
    ("h2", "1. 需要解决的问题"),
    ("p", "1. 高额出售没有代价：玩家把价格定得再高，结算仍然按参考价，定价失去意义。"),
    ("p", "2. 买家个性没有体现：所有买家的提问、语气和反应都一样。"),
    ("p", "3. 选项无脑选第一个：每题都存在唯一最优答案，成交没有紧张感。"),
    ("p", "4. 提问和物品状态脱钩：翻来覆去只有三个问题，回答也不看物品真实状态。"),
    ("p", "5. 说谎没有成本：骗买家不会被识破，信誉也不会掉。"),
    ("h2", "2. 高额出售的代价"),
    ("p", "每个买家心里有一个感知价：感知价 = 物品价值 × 市场趋势 × 藏品加成 × 买家溢价系数。"),
    ("p", "• 挂牌价高于感知价的部分不会成交：买家先压价到感知价，压不下来就离店。"),
    ("p", "• 挂牌价越高，买家初始信任越低：超出参考售价上限，按超出比例每 10% 扣 4 点信任，最多扣 20 点。"),
    ("p", "• 超出参考上限 50% 以上：买家有 30% 概率直接离店，物品回库，下一次到店时间推迟一天。"),
    ("p", "• 定价低于参考下限：成交更快、信任 +3，但单次利润更低。"),
    ("p", "• 买家溢价系数 0.85 至 1.35，由买家性格决定：价敏感型偏低，收藏型偏高。"),
    ("h2", "3. 买家个性"),
    ("p", "每个买家档案带两个轴：核验强度（松 / 中 / 严）和议价风格（爽快 / 磨价 / 挑剔），"
          "再配一个关注点（价格 / 来源 / 品相 / 故事）。"),
    ("p", "• 爽快：只问 1 至 2 个问题，很少压价，信任门槛低。"),
    ("p", "• 磨价：固定压价 10% 至 20%，信任越高压得越少。"),
    ("p", "• 挑剔：问题最多，逐条核对，答错扣的信任更多。"),
    ("p", "• 价敏感：关心价格和同行报价，玩家定价过高时信任掉得最多。"),
    ("p", "• 收藏型：关心来源、年代、成套，真实信息给的奖励高，说谎惩罚也重。"),
    ("p", "买家的开场白、提问用词和压价方式都要体现性格，买家窗口显示一行性格标签。"),
    ("h2", "4. 选项要有取舍"),
    ("p", "1. 每个问题给 3 个选项：如实回答（信任 +，可能暴露瑕疵）、含糊回避（信任 0 至 -1，风险低）、"
          "编造（信任 + 更多，但进入识破判定）。"),
    ("p", "2. 编造选项有基础识破率：核验松 20%、中 30%、严 45%；物品状态和说法矛盾、"
          "玩家没完成相关检视时会额外提高。"),
    ("p", "3. 被识破：信誉 -5、麻烦 +1、买家立刻离店；连续被识破 3 次，该买家 3 天内不再上门。"),
    ("p", "4. 正确答案由物品真实状态决定，选项顺序每次随机，避免玩家连点同一个位置。"),
    ("p", "5. 如实回答不是永远最优：主动说出瑕疵会让买家压价，但信任更高、成交更稳。"),
    ("h2", "5. 问题跟着物品状态走"),
    ("p", "1. 问题从物品实际状态生成：受潮、锈蚀、污渍、破损、编号磨损、来源缺失、维修痕迹等。"),
    ("p", "2. 例：物品受潮时问「有没有受潮？」，如实回答「边角有点潮」信任 +2；"
          "回答「从来没有受潮」进入识破判定，被识破信誉 -5。"),
    ("p", "3. 例：物品编号没核实过时问「编号能对上吗？」，如实回答「我还没核对过」不扣信任，"
          "硬说「肯定对得上」则有识破风险。"),
    ("p", "4. 奢侈品问价格、成套、年代和来源；违禁品问渠道，回答会影响关注度。"),
    ("p", "5. 问题数量：普通物品 2 至 3 个，奢侈品 3 个，挂牌价高于参考上限时额外多问 1 个价格问题。"),
    ("h2", "6. 数值建议"),
    ("p", "• 超出参考上限的信任惩罚：超出比例 × 40，最多 -20。"),
    ("p", "• 编造被识破基础概率：20% / 30% / 45%（对应核验松 / 中 / 严）。"),
    ("p", "• 被识破后果：信誉 -5，麻烦 +1，交易取消。"),
    ("p", "• 压价幅度：10% 至 20%，信任越高压得越少。"),
    ("p", "• 买家溢价系数：0.85 至 1.35。"),
    ("h2", "7. 落地顺序"),
    ("p", "1. 先做挂牌价高于上限的代价和压价，改动最小，能立刻让定价有意义。"),
    ("p", "2. 再做买家性格标签和开场白差异。"),
    ("p", "3. 然后按物品状态生成问题库。"),
    ("p", "4. 最后做编造与识破判定。"),
]


def build_buyer_doc():
    document = docx.Document()
    for kind, text in BUYER_CONTENT:
        if kind == "h1":
            document.add_heading(text, level=1)
        elif kind == "h2":
            document.add_heading(text, level=2)
        else:
            document.add_paragraph(text)
    document.save(BUYER)
    print("created: 买家系统改版方案（待确认）")


def main():
    update_shop()
    update_master()
    build_buyer_doc()
    return 0


if __name__ == "__main__":
    sys.exit(main())
