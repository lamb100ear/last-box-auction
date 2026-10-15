# -*- coding: utf-8 -*-
"""在《买家详细故事文档》的通用规则里补上线索物品的开出与解锁规则。"""

import copy
import os
import shutil
import sys

import docx
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(ROOT, "再拍最后一箱 2.0 买家详细故事文档.docx")

OLD_RULE = "只有线索物品和买家对话产生的关键词可以记录到记事本。"
NEW_RULE = "• 线索物品从拍卖箱随机开出，每件只出现一次，开出即解锁对应买家的故事线。"
EXTRA_RULE = (
    "• 线索物品在检视台处理后，关键词以重点色标出并自动记入记事本；买家对话也会产生关键词。"
)


def set_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def insert_after(anchor, text):
    new = anchor.insert_paragraph_before("", style=anchor.style)
    anchor._element.addnext(new._element)
    run = new.add_run(text)
    rpr = anchor._element.find(qn("w:rPr"))
    if rpr is not None:
        run._element.insert(0, copy.deepcopy(rpr))
    return new


def main():
    document = docx.Document(DOCX)
    for paragraph in document.paragraphs:
        if NEW_RULE in paragraph.text or EXTRA_RULE in paragraph.text:
            print("skip: 通用规则已包含线索物品规则")
            return 0

    backup = os.path.join(os.environ.get("TEMP", "."), "买家详细故事文档.before-rules.docx")
    if not os.path.exists(backup):
        shutil.copy2(DOCX, backup)

    hit = None
    for paragraph in document.paragraphs:
        if OLD_RULE in paragraph.text:
            hit = paragraph
            break
    if hit is None:
        raise SystemExit("找不到通用规则里的关键词条目")

    set_text(hit, NEW_RULE)
    insert_after(hit, EXTRA_RULE)
    document.save(DOCX)
    print("updated: 买家详细故事文档（通用规则补充线索物品开箱规则）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
